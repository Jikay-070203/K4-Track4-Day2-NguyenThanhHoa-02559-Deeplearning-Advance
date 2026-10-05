"""report_tools.py - bảng results.xlsx, biểu đồ tổng hợp và tiện ích viết báo cáo (Bước 5).

Không chứa logic thí nghiệm; chỉ định dạng đầu ra từ các số đã đo. Mọi con số trong bảng đến từ summary.json của
từng run, logit lưu trên đĩa và file dự đoán (tính lại bằng eval.py).
"""
from __future__ import annotations

import math
from pathlib import Path

import numpy as np
import pandas as pd

SHEETS = ["Backbones", "Training", "Inference", "Final", "PerClass", "Latency", "Summary"]


# --------------------------------------------------------------------------- định dạng
def fmt_pm(mean: float, std: float, digits: int = 4) -> str:
    """'0.9412 ± 0.0031' (std NaN khi chỉ có 1 seed -> ghi rõ)."""
    if std is None or (isinstance(std, float) and math.isnan(std)):
        return f"{mean:.{digits}f} (1 seed)"
    return f"{mean:.{digits}f} ± {std:.{digits}f}"


def md_table(df: pd.DataFrame, digits: int = 4) -> str:
    """DataFrame -> bảng markdown (float làm tròn, NaN/None -> '—')."""
    def f(v):
        if v is None or (isinstance(v, float) and math.isnan(v)):
            return "—"
        if isinstance(v, (float, np.floating)):
            return f"{v:.{digits}f}"
        return str(v)
    esc = lambda s: str(s).replace("|", "\\|").replace("\n", " ")        # '|' trong ô sẽ làm vỡ bảng markdown
    cols = list(df.columns)
    lines = ["| " + " | ".join(esc(c) for c in cols) + " |", "|" + "---|" * len(cols)]
    lines += ["| " + " | ".join(esc(f(v)) for v in row) + " |" for row in df.itertuples(index=False)]
    return "\n".join(lines)


# --------------------------------------------------------------------------- results.xlsx
def write_results_xlsx(path: str | Path, sheets: dict[str, pd.DataFrame], best: dict[str, tuple[str, str]] | None = None) -> None:
    """Ghi các sheet, đóng băng hàng tiêu đề, số 4 chữ số, độ rộng cột vừa, tô nổi bật dòng tốt nhất.

    best: {tên sheet: (tên cột, "max" | "min")} -> dòng có giá trị tốt nhất của cột đó được tô xanh.
    """
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        for name in SHEETS:
            df = sheets.get(name)
            if df is None:
                df = pd.DataFrame({"ghi chú": ["(không có dữ liệu)"]})
            df.to_excel(w, sheet_name=name, index=False)
    wb = load_workbook(path)
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    best_fill = PatternFill("solid", fgColor="C6EFCE")
    for name in SHEETS:
        ws = wb[name]
        ws.freeze_panes = "A2"
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(10, width + 2), 60)
            for c in col[1:]:
                if isinstance(c.value, float):
                    c.number_format = "0.0000"
        if best and name in best and name in sheets and len(sheets[name]):
            col_name, mode = best[name]
            df = sheets[name]
            if col_name in df.columns and df[col_name].notna().any():
                idx = df[col_name].astype(float).idxmax() if mode == "max" else df[col_name].astype(float).idxmin()
                for c in ws[idx + 2]:
                    c.fill = best_fill
    wb.save(path)


# --------------------------------------------------------------------------- biểu đồ
def _save(fig, path):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=130, bbox_inches="tight")


def plot_class_distribution(per_class: pd.DataFrame, path: str | Path) -> None:
    """Cột nhóm số ảnh mỗi lớp trong train / val / test. per_class: index = tên lớp, cột train/val/test."""
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(1, 2, figsize=(14, 4.4))
    x = np.arange(len(per_class))
    for k, col in enumerate(["train", "val", "test"]):
        ax[0].bar(x + (k - 1) * 0.27, per_class[col].values, width=0.27, label=col)
    ax[0].set_xticks(x); ax[0].set_xticklabels(per_class.index, rotation=35, ha="right")
    ax[0].set_ylabel("số ảnh"); ax[0].set_title("Số ảnh mỗi lớp theo tập (fold 0)"); ax[0].legend(); ax[0].grid(alpha=0.3, axis="y")
    frac = per_class[["train", "val", "test"]].div(per_class[["train", "val", "test"]].sum(0), axis=1) * 100
    for k, col in enumerate(["train", "val", "test"]):
        ax[1].bar(x + (k - 1) * 0.27, frac[col].values, width=0.27, label=col)
    ax[1].set_xticks(x); ax[1].set_xticklabels(per_class.index, rotation=35, ha="right")
    ax[1].set_ylabel("% ảnh của tập"); ax[1].set_title("Tỉ lệ lớp trong từng tập"); ax[1].grid(alpha=0.3, axis="y")
    fig.tight_layout()
    _save(fig, path)


def plot_tradeoff(df: pd.DataFrame, x: str, y: str, label: str, path: str | Path, xlabel: str, ylabel: str,
                  title: str, logx: bool = False, highlight: list[str] | None = None) -> None:
    """Scatter đánh đổi (ví dụ macro-F1 val theo độ trễ p50). Mỗi điểm ghi nhãn."""
    import matplotlib.pyplot as plt
    d = df.dropna(subset=[x, y])
    fig, ax = plt.subplots(figsize=(8.5, 5.2))
    ax.scatter(d[x], d[y], s=45, c="C0")
    for _, r in d.iterrows():
        ax.annotate(str(r[label]), (r[x], r[y]), textcoords="offset points", xytext=(5, 4), fontsize=8,
                    color="C3" if highlight and r[label] in highlight else "black")
    if logx:
        ax.set_xscale("log")
    ax.set_xlabel(xlabel); ax.set_ylabel(ylabel); ax.set_title(title); ax.grid(alpha=0.3)
    fig.tight_layout()
    _save(fig, path)


def plot_ablation_bars(rows: pd.DataFrame, ref: float, thr: float, path: str | Path, title: str) -> None:
    """Thanh ngang Δ macro-F1 val so với T00 cho từng thí nghiệm; đường đứt = ±ngưỡng nhiễu (2σ)."""
    import matplotlib.pyplot as plt
    d = rows.sort_values("delta")
    fig, ax = plt.subplots(figsize=(8.5, 0.38 * len(d) + 1.8))
    colors = ["C2" if v > thr else ("C3" if v < -thr else "C7") for v in d["delta"]]
    ax.barh(d["label"], d["delta"], color=colors)
    ax.axvline(0, color="k", lw=0.8)
    ax.axvline(thr, color="gray", ls="--", lw=1, label=f"±2σ = {thr:.4f}")
    ax.axvline(-thr, color="gray", ls="--", lw=1)
    ax.set_xlabel(f"Δ macro-F1 val so với T00 (= {ref:.4f})"); ax.set_title(title); ax.legend(loc="lower right")
    ax.grid(alpha=0.3, axis="x")
    fig.tight_layout()
    _save(fig, path)


def plot_confusion(cm: np.ndarray, names: list[str], path: str | Path, title: str) -> None:
    """Ma trận nhầm lẫn: số ảnh (chữ) và recall theo hàng (màu)."""
    import matplotlib.pyplot as plt
    cm = np.asarray(cm)
    norm = cm / cm.sum(1, keepdims=True).clip(min=1)
    fig, ax = plt.subplots(figsize=(8.2, 7))
    im = ax.imshow(norm, cmap="Blues", vmin=0, vmax=1)
    for i in range(cm.shape[0]):
        for j in range(cm.shape[1]):
            ax.text(j, i, f"{int(cm[i, j])}", ha="center", va="center", fontsize=8,
                    color="white" if norm[i, j] > 0.5 else "black")
    ax.set_xticks(range(len(names))); ax.set_yticks(range(len(names)))
    ax.set_xticklabels(names, rotation=40, ha="right"); ax.set_yticklabels(names)
    ax.set_xlabel("dự đoán"); ax.set_ylabel("nhãn thật"); ax.set_title(title)
    fig.colorbar(im, ax=ax, label="tỉ lệ theo hàng (recall)")
    fig.tight_layout()
    _save(fig, path)
